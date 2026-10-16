"""Build an editable Word handout from the released HW2B Markdown."""
import argparse,re
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Polygon,Rectangle

root=Path(__file__).resolve().parents[1]
source=root/'assignments/homework-02/HW2B'
ap=argparse.ArgumentParser();ap.add_argument('--out',type=Path,required=True);args=ap.parse_args()
out=args.out;out.mkdir(parents=True,exist_ok=True)
stem='MIE446_HW2B_Lilienthal_Wing_Beam_Analysis'
doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 for e in list(st.element.iter()):
  if e.tag==qn('w:pBdr'):e.getparent().remove(e)
  if e.tag==qn('w:rFonts'):
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:e.attrib.pop(qn('w:'+attr),None)
  if e.tag==qn('w:color'):
   for attr in ['themeColor','themeTint','themeShade']:e.attrib.pop(qn('w:'+attr),None)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.core_properties.title='HW 2B From Lilienthal wing control loads to beam bending'
doc.core_properties.author='Dr. Ehsan Roohi Golkhatmi'

def mr(s):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def group(tag,children):
 e=OxmlElement('m:'+tag)
 for c in children:e.append(deepcopy(c))
 return e
def fraction(n,d):
 return group('f',[group('num',n),group('den',d)])
def integral(expr):
 n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
 ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∫');pr.append(ch);n.append(pr)
 for tag,text in [('sub','0'),('sup','L')]:n.append(group(tag,[mr(text)]))
 n.append(group('e',[mr(expr)]));return n
def equation(parts):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
 om=OxmlElement('m:oMath')
 for part in parts:om.append(part if not isinstance(part,str) else mr(part))
 p._p.append(om)
 return p
def equations(n):
 if n==0:
  equation(['θ = ',fraction([mr('dv')],[mr('dx')]),'     EI ',fraction([mr('d²v')],[mr('dx²')]),' = −M(x)']).paragraph_format.keep_with_next=True
  equation([group('sSub',[group('e',[mr('σ')]),group('sub',[mr('xx')])]),' = −',fraction([mr('M(x)z')],[mr('I')])])
 elif n==1:
  for name,expr in [('θ(L)','M(x) dx'),('v(L)','(L − x) M(x) dx')]:
   equation([name+' = −',fraction([mr('1')],[mr('EI')]),integral(expr),' = 0']).paragraph_format.keep_with_next=(name=='θ(L)')
 else:
  equation(['Δ',group('sSub',[group('e',[mr('f')]),group('sub',[mr('z')])]),'(s) = q c(s) ',group('sSub',[group('e',[mr('A')]),group('sub',[mr('z')])]),' g(s)     q = ',fraction([mr('1')],[mr('2')]),'ρV²'])

def inline(p,s):
 # Preserve emphasis and external/relative hyperlinks without visible Markdown.
 for token in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*)',s):
  m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
  if m:
   label,url=m.groups()
   if not url.startswith('http'):
    target=(source/url).resolve().relative_to(root).as_posix()
    url='https://github.com/Ehsan-Roohi/Aerospace-Structures/blob/main/'+target
   link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
   r=OxmlElement('w:r');pr=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
   t=OxmlElement('w:t');t.text=label;r.append(t);link.append(r);p._p.append(link)
  else:
   r=p.add_run(token[2:-2] if token.startswith('**') else token)
   r.bold=token.startswith('**')

def figure():
 fig=plt.figure(figsize=(10,5.6));gs=fig.add_gridspec(2,3,height_ratios=[1.3,1])
 ax=fig.add_subplot(gs[0,:]);ax.set(xlim=(-.1,1.1),ylim=(-.4,.6));ax.axis('off')
 ax.add_patch(Polygon([[0,-.23],[1,-.08],[1,.16],[0,.27]],fc='#e5eff3',ec='#536c80',lw=2))
 for x in [.2,.4,.6,.8]:ax.plot([x,x],[-.23+.15*x,.27-.11*x],color='#9db1b8')
 ax.plot([0,1],[0,0],color='#a5422c',lw=5)
 ax.annotate('',xy=(.5,.02),xytext=(.5,.43),arrowprops=dict(arrowstyle='-|>',color='#8a237f',lw=2))
 ax.text(.53,.38,'Prescribed load increment\n3 N at 225 mm',fontsize=12)
 ax.text(.5,-.37,'Scaled member L = 450 mm; schematic, not a historical reconstruction',ha='center',fontsize=12)
 for j,(title,ends) in enumerate([('A  Cantilever',['fixed',None]),('B  Pin and roller',['pin','roller']),('C  Fixed fixed',['fixed','fixed'])]):
  a=fig.add_subplot(gs[1,j]);a.set(xlim=(-.15,1.15),ylim=(-.8,.6));a.axis('off')
  a.plot([0,1],[0,0],lw=4,color='#18334d')
  for x,kind in zip([0,1],ends):
   if kind=='fixed':a.add_patch(Rectangle((x-.035,-.2),.07,.4,fc='#d2e5d7',ec='#277447',lw=2))
   elif kind:
    a.add_patch(Polygon([[x,-.03],[x-.09,-.23],[x+.09,-.23]],fc='#d2e5d7',ec='#277447',lw=2))
    if kind=='roller':a.plot([x-.04,x+.04],[-.29,-.29],'o',ms=4,color='#277447')
  a.annotate('',xy=(.5,.02),xytext=(.5,.45),arrowprops=dict(arrowstyle='-|>',color='#8a237f',lw=2))
  a.text(.5,-.52,title,ha='center',fontsize=12,fontweight='bold')
 fig.suptitle('One wing member with three alternative physical fixtures',fontsize=15)
 fig.text(.5,.02,'Applied load and supports only. Students draw the separate reaction free bodies.',ha='center',fontsize=11)
 fig.tight_layout(rect=[0,.07,1,.95]);path=out/'member_models.png';fig.savefig(path,dpi=220);plt.close(fig)
 p=doc.add_paragraph();p.alignment=1;p.add_run().add_picture(str(path),width=Inches(6.9))
 doc.add_paragraph('Figure 1. Conceptual load path and fixtures. These are alternative teaching models, not three historical wing attachments.',style='Caption')

text=(source/'README.md').read_text(encoding='utf-8')
blocks=re.split(r'\n\s*\n',text.strip());eq=0
for block in blocks:
 if block.startswith('**Download the handout:**'):continue
 if block.startswith('# '):
  doc.add_paragraph(block[2:],style='Title');continue
 if block.startswith('### '):doc.add_heading(block[4:],level=2);continue
 if block.startswith('## '):doc.add_heading(block[3:],level=1);continue
 if block.startswith('!['):figure();continue
 if block.startswith('$$'):equations(eq);eq+=1;continue
 if block.startswith('|'):
  rows=[r.strip().strip('|').split('|') for r in block.splitlines()]
  rows=[r for r in rows if not all(re.fullmatch(r'[:\-\s]+',c) for c in r)]
  table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
  table.columns[0].width=Inches(1.4);table.columns[1].width=Inches(5.5)
  for ri,row in enumerate(rows):
   cells=table.add_row().cells
   for c,val in zip(cells,row):
    inline(c.paragraphs[0],val.strip())
    for p in c.paragraphs:p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tcpr.append(borders)
    if ri==0:
     for p in c.paragraphs:p.paragraph_format.keep_with_next=True
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF0');tcpr.append(shade)
     for run in c.paragraphs[0].runs:run.bold=True
   trpr=table.rows[-1]._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
   if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  continue
 if re.match(r'(- |\d+\. )',block):
  for line in block.splitlines():
   p=doc.add_paragraph()
   p.paragraph_format.left_indent=Inches(.18)
   p.paragraph_format.first_line_indent=Inches(-.12)
   p.paragraph_format.space_after=Pt(4)
   inline(p,('• '+line[2:]) if line.startswith('- ') else line)
 else:inline(doc.add_paragraph(),block.replace('\n',' '))
assert eq==3
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('MIE 446  •  HW 2B  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(out/(stem+'.docx'))
print(out/(stem+'.docx'))
